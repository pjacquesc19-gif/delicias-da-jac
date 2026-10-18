#!/usr/bin/env python3
"""Testes do conversor planilha -> JSON.  Uso: python scripts/test_conversor.py"""
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / 'scripts' / 'xlsx_to_json.py'
XLSX = ROOT / 'data' / 'catalogo.xlsx'
JSON_PATH = ROOT / 'data' / 'produtos.json'

# Catálogo 2026 original: nome -> (preço, pedido mínimo na unidade de venda).
# Mini suspiros: mínimo de 500 g = 2 pacotes de 250 g.
ORIGINAL = {
    'Macaron individual': (6, 20), 'Caixinha com 1 macaron': (8, None), 'Caixinha com 2 macarons': (15, None),
    'Caixinha com 3 macarons': (22, None),
    'Torre de macarons 4 andares - sem torre': (240, None), 'Torre de macarons 4 andares - com torre': (274, None),
    'Torre de macarons 5 andares - sem torre': (360, None), 'Torre de macarons 5 andares - com torre': (395, None),
    'Torre de macarons 6 andares - sem torre': (480, None), 'Torre de macarons 6 andares - com torre': (540, None),
    'Torre de macarons 7 andares - sem torre': (600, None), 'Torre de macarons 7 andares - com torre': (680, None),
    'Torre de macarons 8 andares - sem torre': (780, None), 'Torre de macarons 8 andares - com torre': (880, None),
    'Torre de macarons 10 andares - sem torre': (1100, None), 'Torre de macarons 10 andares - com torre': (1200, None),
    'Pirulito simples com laço': (5.5, 10), 'Pirulito modelado simples': (6, 10), 'Pirulito modelado premium': (7.5, 10),
    'Pirulito recheado de brigadeiro': (6, 10),
    'Saquinho com 15 mini suspiros': (5, 10), 'Saquinho com 15 mini suspiros - cores fortes': (5.5, 10),
    'Rosa recheada no saquinho': (5.5, 10), 'Caixinha com mini suspiros': (7, 10), 'Rosa recheada na caixinha': (7, 10),
    'Bonequinha recheada': (5.5, 10), 'Adicional suspiro temático central': (0.5, None),
    'Suspiro modelado sem recheio': (7.5, 10), 'Suspiro modelado recheado': (8, 10),
    'Suspiro 3D sem caixinha': (10, 10), 'Suspiro 3D com caixinha': (12, 10),
    'Centro de mesa com pirulitos de roseta': (10, 10), 'Centro de mesa com pirulitos modelados': (12, 10),
    'Centro de mesa com pirulitos premium': (14, 10),
    'Torre de suspiros 30 cm': (100, None), 'Torre de suspiros 40 cm': (150, None), 'Mini suspiros': (45, 2),
    'Brownie 4 cm sem laço': (5, 10), 'Brownie 4 cm com laço': (5.3, 10),
    'Brownie 5 cm sem laço': (6, None), 'Brownie 5 cm com laço': (6.3, None),
    'Bem-casado / bem-vivido sem laço': (5.5, 15), 'Bem-casado / bem-vivido com laço': (6, 15),
}


def run(*args, xlsx=None, out=None):
    cmd = [sys.executable, '-I', str(SCRIPT)] + list(args)
    if xlsx:
        cmd += ['--xlsx', str(xlsx)]
    if out:
        cmd += ['--out', str(out)]
    return subprocess.run(cmd, capture_output=True, text=True)


def items(data):
    for p in data['products']:
        if p['variants']:
            for v in p['variants']:
                yield p, v['label'], v['price'], v['minimumQuantity']
        else:
            yield p, p['name'], p['price'], p['minimumQuantity']


class Conversor(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def mutate(self, fn):
        """Copia a planilha, aplica fn(workbook) e devolve o caminho."""
        wb = load_workbook(XLSX)
        fn(wb)
        path = self.tmp / 'teste.xlsx'
        wb.save(path)
        return path

    def expect_error(self, fn, fragment):
        res = run(xlsx=self.mutate(fn), out=self.tmp / 'saida.json')
        self.assertEqual(res.returncode, 1, res.stdout)
        self.assertIn(fragment, res.stdout)
        self.assertFalse((self.tmp / 'saida.json').exists(), 'não deve gravar JSON quando há erros')

    # --- migração e geração ---
    def test_json_do_repositorio_esta_em_dia(self):
        res = run('--check')
        self.assertEqual(res.returncode, 0, res.stdout)

    def test_geracao_e_deterministica(self):
        a, b = self.tmp / 'a.json', self.tmp / 'b.json'
        self.assertEqual(run(out=a).returncode, 0)
        self.assertEqual(run(out=b).returncode, 0)
        self.assertEqual(a.read_bytes(), b.read_bytes())
        self.assertEqual(a.read_bytes(), JSON_PATH.read_bytes())

    def test_precos_e_minimos_batem_com_catalogo_original(self):
        data = json.loads(JSON_PATH.read_text(encoding='utf-8'))
        got = {label: (price, minimum) for _, label, price, minimum in items(data)}
        self.assertEqual(set(got), set(ORIGINAL), 'itens faltando ou sobrando')
        for label, expected in ORIGINAL.items():
            self.assertEqual(got[label], expected, label)
        self.assertEqual(len(ORIGINAL), 43)

    def test_categorias_e_limites_do_catalogo(self):
        data = json.loads(JSON_PATH.read_text(encoding='utf-8'))
        self.assertEqual([c['name'] for c in data['categories']], [
            'Macarons', 'Pirulitos de Suspiro', 'Lembrancinhas', 'Suspiros Modelados', 'Centros de Mesa',
            'Torres e Mini Suspiros', 'Brownies', 'Bem-casados e Bem-vividos'])
        prods = {p['id']: p for p in data['products']}
        torre = {v['label']: v for v in prods['macaron-torre']['variants']}
        self.assertEqual(torre['Torre de macarons 4 andares - com torre']['maxFlavors'], 2)
        self.assertEqual(torre['Torre de macarons 6 andares - sem torre']['maxFlavors'], 3)
        self.assertIsNone(torre['Torre de macarons 10 andares - com torre']['maxFlavors'])
        t30 = {v['label']: v for v in prods['torre-de-suspiros']['variants']}
        self.assertEqual(t30['Torre de suspiros 30 cm']['maxColors'], 2)
        self.assertEqual(prods['brownie']['flavors'], ['Brigadeiro', 'Doce de leite'])
        self.assertEqual(prods['suspiro-3d']['flavors'], ['Nutella', 'Brigadeiro'])
        self.assertEqual(len(prods['macaron-individual']['flavors']), 8)
        self.assertEqual(prods['mini-suspiros']['gramsPerUnit'], 250)
        self.assertEqual(prods['macaron-individual']['minPerColor'], 20)
        self.assertEqual(prods['brownie']['colors'], [])
        self.assertTrue(all(len(p['id']) for p in data['products']))
        self.assertEqual(len({p['id'] for p in data['products']}), len(data['products']))
        consult = [c['name'] for c in data['colors'] if c['consult']]
        self.assertEqual(consult, ['Dourado', 'Prateado'])

    # --- validação ---
    def test_preco_ausente_e_recusado(self):
        def f(wb):
            ws = wb['Produtos']
            for row in ws.iter_rows(min_row=2):
                if row[0].value == 'pirulito-simples':
                    row[4].value = None
        self.expect_error(f, 'preço ausente')

    def test_preco_ausente_em_variacao_e_recusado(self):
        def f(wb):
            ws = wb['Variações']
            ws.cell(2, 7).value = None
        self.expect_error(f, 'preço ausente')

    def test_imagem_com_endereco_externo_e_recusada(self):
        def f(wb):
            ws = wb['Produtos']
            ws.cell(2, 12).value = 'https://site.com/foto.jpg'
        self.expect_error(f, 'nome de imagem inválido')

    def test_imagem_com_pasta_ou_extensao_indevida_e_recusada(self):
        for bad in ('../segredo.jpg', 'pasta/foto.png', 'script.js', 'foto.gif'):
            def f(wb, bad=bad):
                wb['Produtos'].cell(2, 12).value = bad
            self.expect_error(f, 'nome de imagem inválido')

    def test_imagens_validas_sao_aceitas_e_aviso_se_nao_existem(self):
        def f(wb):
            wb['Produtos'].cell(2, 12).value = 'macaron-rosa.jpg; macaron-azul.WEBP'
        res = run(xlsx=self.mutate(f), out=self.tmp / 'saida.json')
        self.assertEqual(res.returncode, 0, res.stdout)
        self.assertIn('AVISO', res.stdout)
        data = json.loads((self.tmp / 'saida.json').read_text(encoding='utf-8'))
        self.assertEqual(data['products'][0]['images'], ['macaron-rosa.jpg', 'macaron-azul.WEBP'])

    def test_categoria_inexistente(self):
        self.expect_error(lambda wb: setattr(wb['Produtos'].cell(2, 2), 'value', 'Salgados'), 'não existe na aba Categorias')

    def test_id_duplicado(self):
        self.expect_error(lambda wb: setattr(wb['Produtos'].cell(3, 1), 'value', 'macaron-individual'), 'repetido')

    def test_pedido_minimo_invalido(self):
        self.expect_error(lambda wb: setattr(wb['Produtos'].cell(5, 7), 'value', 0), 'inteiro')

    def test_variacao_de_produto_inexistente(self):
        self.expect_error(lambda wb: setattr(wb['Variações'].cell(2, 1), 'value', 'nao-existe'), 'não existe na aba Produtos')

    def test_cor_fora_da_lista(self):
        self.expect_error(lambda wb: setattr(wb['Produtos'].cell(5, 16), 'value', 'Rosa; Turquesa'), 'não existe na aba Cores')

    def test_booleano_invalido(self):
        self.expect_error(lambda wb: setattr(wb['Produtos'].cell(2, 22), 'value', 'talvez'), 'Sim ou Não')

    def test_coluna_obrigatoria_removida(self):
        self.expect_error(lambda wb: setattr(wb['Produtos'].cell(1, 3), 'value', 'Título'), 'coluna obrigatória')

    def test_coluna_de_preco_removida_recusa_produtos_sem_variacao(self):
        self.expect_error(lambda wb: setattr(wb['Produtos'].cell(1, 5), 'value', 'Valor'), 'preço ausente')

    def test_ordem_das_colunas_nao_importa(self):
        def f(wb):
            ws = wb['Produtos']
            ws.move_range('C1:D80', rows=0, cols=0)
            ws.insert_cols(1)  # desloca tudo: o conversor localiza colunas pelo nome
        res = run(xlsx=self.mutate(f), out=self.tmp / 'saida.json')
        self.assertEqual(res.returncode, 0, res.stdout)
        self.assertEqual((self.tmp / 'saida.json').read_bytes(), JSON_PATH.read_bytes())

    def test_preco_com_virgula_em_texto(self):
        def f(wb):
            ws = wb['Variações']
            for row in ws.iter_rows(min_row=2):
                if row[1].value == 'Brownie 4 cm com laço':
                    row[6].value = 'R$ 5,30'
        res = run(xlsx=self.mutate(f), out=self.tmp / 'saida.json')
        self.assertEqual(res.returncode, 0, res.stdout)
        self.assertEqual((self.tmp / 'saida.json').read_bytes(), JSON_PATH.read_bytes())


if __name__ == '__main__':
    unittest.main(verbosity=2)
