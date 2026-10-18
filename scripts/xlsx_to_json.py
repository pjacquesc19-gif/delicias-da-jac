#!/usr/bin/env python3
"""Converte data/catalogo.xlsx em data/produtos.json, validando os dados.

Uso:
    python scripts/xlsx_to_json.py            # gera data/produtos.json
    python scripts/xlsx_to_json.py --check    # só confere se o JSON está em dia

Sai com código 1 (e não grava nada) se houver erros na planilha.
O JSON é determinístico (sem data/hora): se a planilha não mudou, o arquivo
gerado é idêntico ao anterior, o que evita commits desnecessários.
"""
import argparse
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
ERRORS, WARNINGS = [], []

IMG_RE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9._-]*\.(jpg|jpeg|png|webp)$', re.I)
BRAND_IMG_RE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9._-]*\.(jpg|jpeg|png|webp|svg)$', re.I)
SLUG_RE = re.compile(r'^[a-z0-9]+(-[a-z0-9]+)*$')
HEX_RE = re.compile(r'^#[0-9A-Fa-f]{6}$')


# ---------- utilidades ----------
def norm(s):
    s = unicodedata.normalize('NFD', str(s))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', s).strip().lower()


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', norm(s)).strip('-')


def err(where, msg):
    ERRORS.append(f'[{where}] {msg}')


def warn(where, msg):
    WARNINGS.append(f'[{where}] {msg}')


def text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_bool(v, where, default):
    t = norm(text(v))
    if t == '':
        return default
    if t in ('sim', 's', 'true', '1', 'x', 'yes'):
        return True
    if t in ('nao', 'n', 'false', '0', 'no'):
        return False
    err(where, f'valor "{v}" inválido: use Sim ou Não.')
    return default


def to_number(v, where, label):
    if v is None or text(v) == '':
        return None
    if isinstance(v, (int, float)):
        n = float(v)
    else:
        t = text(v).replace('R$', '').replace(' ', '')
        if ',' in t:
            t = t.replace('.', '').replace(',', '.')
        try:
            n = float(t)
        except ValueError:
            err(where, f'{label}: "{v}" não é um número.')
            return None
    if n < 0:
        err(where, f'{label}: não pode ser negativo.')
        return None
    return round(n, 2)


def to_int(v, where, label, minimum=1):
    n = to_number(v, where, label)
    if n is None:
        return None
    if not float(n).is_integer() or n < minimum:
        err(where, f'{label}: use um número inteiro maior ou igual a {minimum}.')
        return None
    return int(n)


def split_list(v):
    out = []
    for part in re.split(r'[;\n]', text(v)):
        part = part.strip()
        if part and part not in out:
            out.append(part)
    return out


# ---------- leitura da planilha ----------
def read_sheet(wb, name, aliases, required):
    """Lê uma aba pelos nomes das colunas (a ordem das colunas não importa)."""
    sheet = next((wb[n] for n in wb.sheetnames if norm(n) == norm(name)), None)
    if sheet is None:
        err(name, 'aba não encontrada.')
        return []
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        err(name, 'aba vazia.')
        return []
    idx = {}
    for i, h in enumerate(rows[0]):
        key = aliases.get(norm(h)) if h is not None else None
        if key:
            idx[key] = i
    for key in required:
        if key not in idx:
            label = next(h for h, k in aliases.items() if k == key)
            err(name, f'coluna obrigatória "{label}" não encontrada na primeira linha.')
    out = []
    for n, row in enumerate(rows[1:], start=2):
        rec = {k: None for k in set(aliases.values())}  # colunas ausentes viram vazio (o erro já foi registrado)
        rec.update({k: (row[i] if i < len(row) else None) for k, i in idx.items()})
        if any(text(v) for v in rec.values()):
            rec['_row'] = n
            out.append(rec)
    return out


PROD_COLS = {
    'id': 'id', 'categoria': 'cat', 'nome': 'nome', 'descricao': 'desc', 'preco (r$)': 'preco',
    'texto do preco': 'texto_preco', 'pedido minimo': 'minimo', 'quantidade / tamanho': 'qtd',
    'unidade (singular)': 'un_s', 'unidade (plural)': 'un_p', 'gramas por unidade': 'gramas',
    'imagens': 'imagens', 'sabores': 'sabores', 'escolha de sabores': 'escolha', 'minimo por cor': 'min_cor',
    'cores': 'cores', 'maximo de cores': 'max_cores', 'temas': 'temas', 'personalizavel': 'pers',
    'etiqueta': 'etiqueta', 'observacoes': 'obs', 'destaque': 'destaque', 'disponivel': 'disp',
}
VAR_COLS = {
    'id do produto': 'pid', 'variacao': 'label', 'seletor 1': 's1', 'valor 1': 'v1', 'seletor 2': 's2',
    'valor 2': 'v2', 'preco (r$)': 'preco', 'pedido minimo': 'minimo', 'quantidade / tamanho': 'qtd',
    'maximo de sabores': 'max_sab', 'maximo de cores': 'max_cores', 'sabores (substitui)': 'sabores',
    'temas (substitui)': 'temas', 'observacoes': 'obs',
}


def valid_image(name, where, brand=False):
    ok = (BRAND_IMG_RE if brand else IMG_RE).match(name) and '..' not in name
    if not ok:
        err(where, f'nome de imagem inválido "{name}": use só o nome do arquivo (letras, números, ponto, hífen e '
                   f'sublinhado), sem pastas nem endereços, terminando em .jpg, .jpeg, .png ou .webp.')
    return bool(ok)


def parse_max_colors(v, where, default):
    t = norm(text(v))
    if t == '':
        return default
    if t in ('livre', 'ilimitado', 'sem limite'):
        return 'free'
    n = to_int(v, where, 'Máximo de cores')
    return n if n else default


def parse_override(v, expand=None):
    """None = herdar; [] = não oferece ('-'); lista = valores próprios."""
    t = text(v)
    if t == '':
        return None
    if norm(t) in ('-', 'nenhum', 'nenhuma', 'nao'):
        return []
    if expand is not None and norm(t) in ('padrao', 'todas', 'todos'):
        return list(expand)
    return split_list(t)


# ---------- conversão ----------
def convert(xlsx_path, images_dir):
    wb = load_workbook(xlsx_path, data_only=True)

    # Categorias
    categories, cat_by_name = [], {}
    for r in read_sheet(wb, 'Categorias', {'id': 'id', 'nome': 'nome', 'descricao': 'desc'}, ['id', 'nome']):
        w = f'Categorias, linha {r["_row"]}'
        cid, name = text(r['id']), text(r['nome'])
        if not SLUG_RE.match(cid):
            err(w, f'ID "{cid}" inválido: use letras minúsculas, números e hífen.')
        if norm(name) in cat_by_name:
            err(w, f'categoria "{name}" repetida.')
        cat_by_name[norm(name)] = cid
        categories.append({'id': cid, 'name': name, 'description': text(r.get('desc'))})
    cat_ids = {c['id'] for c in categories}

    # Cores
    colors = []
    for r in read_sheet(wb, 'Cores', {'nome': 'nome', 'cor (hex)': 'hex', 'sob consulta': 'consulta'}, ['nome', 'hex']):
        w = f'Cores, linha {r["_row"]}'
        hx = text(r['hex'])
        if not HEX_RE.match(hx):
            err(w, f'cor "{hx}" inválida: use o formato #RRGGBB.')
        if norm(text(r['nome'])) == 'outro':
            warn(w, 'a opção "Outro" é adicionada automaticamente; remova esta linha.')
            continue
        colors.append({'name': text(r['nome']), 'hex': hx.upper(), 'consult': to_bool(r.get('consulta'), w, False)})
    color_names = [c['name'] for c in colors]

    # Temas
    themes = [text(r['tema']) for r in read_sheet(wb, 'Temas', {'tema': 'tema'}, ['tema'])
              if text(r['tema']) and norm(text(r['tema'])) not in ('outros', 'outro')]

    # Informações gerais e Sobre
    info = {norm(text(r['campo'])): text(r['info']) for r in
            read_sheet(wb, 'Informações gerais', {'campo': 'campo', 'informacao': 'info'}, ['campo', 'info'])}
    about = [text(r['p']) for r in read_sheet(wb, 'Sobre', {'paragrafo': 'p'}, ['p']) if text(r['p'])]
    whatsapp = re.sub(r'\D', '', info.get('whatsapp (numero com ddi)', ''))
    if not re.fullmatch(r'55\d{10,11}', whatsapp):
        err('Informações gerais', 'WhatsApp (número com DDI) deve ter só números, no formato 55 + DDD + número.')
    logo, owner_photo = info.get('logomarca (arquivo)', ''), info.get('foto da proprietaria (arquivo)', '')
    for fname in (logo, owner_photo):
        if fname:
            valid_image(fname, 'Informações gerais', brand=True)
    endpoint = info.get('sistema de encomendas (url do apps script)', '').strip()
    if endpoint and not re.fullmatch(r'https://script\.google\.com/macros/s/[A-Za-z0-9_-]+/exec', endpoint):
        err('Informações gerais', 'endereço do sistema de encomendas inválido: use o endereço do aplicativo web '
            '(https://script.google.com/macros/s/.../exec) ou deixe vazio.')
    meta = {
        'ordersEndpoint': endpoint or None,
        'brand': info.get('marca', ''), 'whatsapp': whatsapp, 'whatsappLabel': info.get('contato', ''),
        'instagram': info.get('instagram', '').lstrip('@'), 'owner': info.get('proprietaria', ''),
        'logo': logo or None, 'ownerPhoto': owner_photo or None, 'about': about,
    }
    if not meta['brand']:
        err('Informações gerais', 'campo "Marca" vazio.')

    # Variações agrupadas por produto
    variations = {}
    for r in read_sheet(wb, 'Variações', VAR_COLS, ['pid', 'label', 'preco']):
        variations.setdefault(text(r['pid']), []).append(r)

    # Produtos
    products, seen_ids, missing_images = [], set(), set()
    prod_rows = read_sheet(wb, 'Produtos', PROD_COLS, ['id', 'cat', 'nome'])
    for r in prod_rows:
        w = f'Produtos, linha {r["_row"]}'
        pid, name = text(r['id']), text(r['nome'])
        if not SLUG_RE.match(pid):
            err(w, f'ID "{pid}" inválido: use letras minúsculas, números e hífen (ex.: brownie).')
        if pid in seen_ids:
            err(w, f'ID "{pid}" repetido.')
        seen_ids.add(pid)
        if not name:
            err(w, 'nome vazio.')
        cat = cat_by_name.get(norm(text(r['cat'])))
        if cat is None:
            err(w, f'categoria "{text(r["cat"])}" não existe na aba Categorias.')
        w = f'Produto "{pid}"'

        un_s, un_p = text(r.get('un_s')) or 'unidade', text(r.get('un_p')) or 'unidades'
        grams = to_int(r.get('gramas'), w, 'Gramas por unidade')
        personalizable = to_bool(r.get('pers'), w, True)

        imgs = []
        for fname in split_list(r.get('imagens')):
            if valid_image(fname, w):
                imgs.append(fname)
                if images_dir and not (images_dir / fname).exists():
                    missing_images.add(f'{pid}: {fname}')

        flavors = split_list(r.get('sabores'))
        choice = norm(text(r.get('escolha')))
        if choice in ('um', 'uma', '1'):
            flavor_choice = 'one'
        elif choice in ('varios', 'varias', 'mais de um', 'multiplo'):
            flavor_choice = 'many'
        elif choice == '':
            flavor_choice = 'many'
        else:
            err(w, f'Escolha de sabores "{text(r.get("escolha"))}" inválida: use um ou vários.')
            flavor_choice = 'many'

        def expand(cell, universe, label):
            t = norm(text(cell))
            if t == '':
                return []
            if t in ('padrao', 'todas', 'todos'):
                return list(universe)
            return split_list(cell)

        prod_colors = expand(r.get('cores'), color_names, 'Cores') if personalizable else []
        prod_themes = expand(r.get('temas'), themes, 'Temas') if personalizable else []
        if not personalizable and (text(r.get('cores')) or text(r.get('temas'))):
            warn(w, 'Personalizável = Não: Cores e Temas foram ignorados.')
        for c in prod_colors:
            if c not in color_names:
                err(w, f'cor "{c}" não existe na aba Cores.')

        notes = split_list(r.get('obs'))
        max_colors = parse_max_colors(r.get('max_cores'), w, 1)
        min_per_color = to_int(r.get('min_cor'), w, 'Mínimo por cor')
        if min_per_color and not prod_colors:
            warn(w, '"Mínimo por cor" ignorado porque o produto não oferece cores.')
            min_per_color = None

        prod = {
            'id': pid, 'category': cat, 'name': name, 'description': text(r.get('desc')),
            'price': None, 'priceUnit': text(r.get('texto_preco')) or f'por {un_s}',
            'unit': [un_s, un_p], 'gramsPerUnit': grams,
            'minimumQuantity': None, 'quantity': text(r.get('qtd')) or None,
            'images': imgs, 'badge': text(r.get('etiqueta')) or None,
            'flavors': flavors, 'flavorChoice': flavor_choice,
            'minPerColor': min_per_color, 'colors': prod_colors, 'maxColors': max_colors,
            'themes': prod_themes, 'customizable': personalizable, 'notes': notes,
            'featured': to_bool(r.get('destaque'), w, False), 'available': to_bool(r.get('disp'), w, True),
            'selectors': [], 'variants': [],
        }

        vrows = variations.pop(pid, [])
        price = to_number(r.get('preco'), w, 'Preço')
        minimum = to_int(r.get('minimo'), w, 'Pedido mínimo')
        if vrows:
            if price is not None:
                err(w, 'tem preço na aba Produtos e também variações: deixe o preço só nas variações.')
            if minimum is not None:
                warn(w, 'Pedido mínimo da aba Produtos ignorado: em produtos com variações, use a aba Variações.')
            sel_labels, labels_seen, combos = [], set(), set()
            for vr in vrows:
                vw = f'Variações, linha {vr["_row"]} ({pid})'
                pairs = [(text(vr.get('s1')), text(vr.get('v1'))), (text(vr.get('s2')), text(vr.get('v2')))]
                pairs = [p for p in pairs if p[0] or p[1]]
                if not pairs:
                    err(vw, 'preencha Seletor 1 e Valor 1 (o que o cliente escolhe, ex.: Tamanho / 4 cm).')
                    continue
                for s, v in pairs:
                    if not s or not v:
                        err(vw, 'cada seletor precisa de nome e valor.')
                labels = [p[0] for p in pairs]
                if not sel_labels:
                    sel_labels = labels
                elif labels != sel_labels:
                    err(vw, f'seletores {labels} diferentes dos das outras variações ({sel_labels}).')
                label = text(vr.get('label'))
                if not label:
                    err(vw, 'nome da variação vazio.')
                if label in labels_seen:
                    err(vw, f'variação "{label}" repetida.')
                labels_seen.add(label)
                combo = tuple(p[1] for p in pairs)
                if combo in combos:
                    err(vw, f'combinação {combo} repetida.')
                combos.add(combo)
                vprice = to_number(vr.get('preco'), vw, 'Preço')
                if vprice is None:
                    err(vw, 'preço ausente. O catálogo não aceita preços vazios.')
                prod['variants'].append({
                    'label': label,
                    'options': {slug(s): v for s, v in pairs},
                    'price': vprice,
                    'minimumQuantity': to_int(vr.get('minimo'), vw, 'Pedido mínimo'),
                    'quantity': text(vr.get('qtd')) or None,
                    'maxFlavors': to_int(vr.get('max_sab'), vw, 'Máximo de sabores'),
                    'maxColors': parse_max_colors(vr.get('max_cores'), vw, None),
                    'flavors': parse_override(vr.get('sabores')),
                    'themes': parse_override(vr.get('temas'), themes) if personalizable else None,
                    'notes': split_list(vr.get('obs')),
                })
            prod['selectors'] = [{'key': slug(s), 'label': s} for s in sel_labels]
        else:
            if price is None:
                err(w, 'preço ausente e sem variações. O catálogo não aceita preços vazios.')
            prod['price'] = price
            prod['minimumQuantity'] = minimum
        products.append(prod)

    for pid in variations:
        err('Variações', f'o ID "{pid}" não existe na aba Produtos.')
    if not products:
        err('Produtos', 'nenhum produto cadastrado.')
    for item in sorted(missing_images):
        warn('Imagens', f'arquivo não encontrado em public/images/produtos/: {item} (o site mostra a imagem padrão).')

    used = {p['category'] for p in products}
    for c in categories:
        if c['id'] not in used:
            warn('Categorias', f'categoria "{c["name"]}" sem produtos: não será exibida.')
    return {'meta': meta, 'categories': categories, 'colors': colors, 'themes': themes, 'products': products}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--xlsx', default=str(ROOT / 'data' / 'catalogo.xlsx'))
    ap.add_argument('--out', default=str(ROOT / 'data' / 'produtos.json'))
    ap.add_argument('--images', default=str(ROOT / 'public' / 'images' / 'produtos'))
    ap.add_argument('--check', action='store_true', help='não grava; falha se o JSON estiver desatualizado')
    args = ap.parse_args()

    data = convert(args.xlsx, Path(args.images))
    gh = os.environ.get('GITHUB_ACTIONS') == 'true'
    for m in WARNINGS:
        print(f'AVISO {m}')
        if gh:
            print(f'::warning::{m}')
    if ERRORS:
        for m in ERRORS:
            print(f'ERRO  {m}')
            if gh:
                print(f'::error::{m}')
        print(f'\n{len(ERRORS)} erro(s) na planilha. O JSON não foi gerado.')
        return 1

    out = json.dumps(data, ensure_ascii=False, indent=2) + '\n'
    out_path = Path(args.out)
    n_var = sum(max(1, len(p['variants'])) for p in data['products'])
    if args.check:
        if not out_path.exists() or out_path.read_text(encoding='utf-8') != out:
            print('data/produtos.json está desatualizado em relação à planilha.')
            return 1
        print('data/produtos.json está em dia.')
        return 0
    out_path.parent.mkdir(parents=True, exist_ok=True)
    changed = not out_path.exists() or out_path.read_text(encoding='utf-8') != out
    if changed:
        out_path.write_text(out, encoding='utf-8')
    print(f'{"Gerado" if changed else "Sem alterações em"} {out_path.name}: {len(data["products"])} produtos, '
          f'{n_var} itens de catálogo, {len(data["categories"])} categorias.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
